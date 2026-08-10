"""Parse security questionnaires from CSV / XLSX."""

from __future__ import annotations

import csv
import re
from pathlib import Path

from openpyxl import load_workbook


QUESTION_HEADERS = {
    "question",
    "questions",
    "control",
    "control question",
    "requirement",
    "security question",
    "item",
    "description",
    "query",
    "问题",
    "控制项",
    "要求",
}

SECTION_HEADERS = {
    "section",
    "category",
    "domain",
    "area",
    "control family",
    "章节",
    "分类",
}

ANSWER_HEADERS = {
    "answer",
    "response",
    "your answer",
    "vendor response",
    "答案",
    "回复",
}


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "").strip().lower())


def _find_col(headers: list[str], candidates: set[str]) -> int | None:
    for i, h in enumerate(headers):
        if _norm(h) in candidates:
            return i
    # fuzzy contains
    for i, h in enumerate(headers):
        nh = _norm(h)
        for c in candidates:
            if c in nh or nh in c:
                return i
    return None


def parse_questionnaire(path: Path) -> list[dict]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        rows = _read_csv(path)
    elif suffix in {".xlsx", ".xlsm"}:
        rows = _read_xlsx(path)
    else:
        raise ValueError("Only .csv and .xlsx questionnaires are supported.")

    if not rows:
        raise ValueError("Empty questionnaire file.")

    # Detect header row
    header_idx = 0
    headers = [str(c) for c in rows[0]]
    q_col = _find_col(headers, QUESTION_HEADERS)
    if q_col is None:
        # Maybe no header — treat first column as question
        return [
            {
                "row_number": i + 1,
                "section": "",
                "question_text": str(row[0]).strip(),
            }
            for i, row in enumerate(rows)
            if row and str(row[0]).strip()
        ]

    s_col = _find_col(headers, SECTION_HEADERS)
    a_col = _find_col(headers, ANSWER_HEADERS)

    questions: list[dict] = []
    for i, row in enumerate(rows[1:], start=2):
        if q_col >= len(row):
            continue
        qtext = str(row[q_col] or "").strip()
        if not qtext or len(qtext) < 3:
            continue
        section = ""
        if s_col is not None and s_col < len(row):
            section = str(row[s_col] or "").strip()
        existing = ""
        if a_col is not None and a_col < len(row):
            existing = str(row[a_col] or "").strip()
        questions.append(
            {
                "row_number": i,
                "section": section,
                "question_text": qtext,
                "existing_answer": existing,
            }
        )
    if not questions:
        raise ValueError("No questions detected. Ensure a 'Question' column exists.")
    return questions


def _read_csv(path: Path) -> list[list]:
    # Try utf-8 then latin-1
    for enc in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            with path.open("r", encoding=enc, newline="") as f:
                reader = csv.reader(f)
                return [list(r) for r in reader]
        except UnicodeDecodeError:
            continue
    raise ValueError("Could not decode CSV file.")


def _read_xlsx(path: Path) -> list[list]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = []
    for row in ws.iter_rows(values_only=True):
        rows.append([("" if c is None else c) for c in row])
    wb.close()
    return rows
