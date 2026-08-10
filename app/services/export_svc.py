"""Export filled questionnaires to CSV / XLSX."""

from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.orm import Session

from app.config import EXPORT_DIR
from app.database import Question, Questionnaire


def export_csv(db: Session, questionnaire_id: int) -> Path:
    qn = db.get(Questionnaire, questionnaire_id)
    if not qn:
        raise ValueError("Not found")
    questions = (
        db.query(Question)
        .filter(Question.questionnaire_id == questionnaire_id)
        .order_by(Question.row_number)
        .all()
    )
    path = EXPORT_DIR / f"questionnaire_{questionnaire_id}_answers.csv"
    with path.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(
            [
                "Row",
                "Section",
                "Question",
                "Answer",
                "Confidence",
                "Status",
                "Sources",
                "Notes",
            ]
        )
        for q in questions:
            sources = ""
            try:
                src = json.loads(q.sources or "[]")
                sources = "; ".join(
                    f"{s.get('title')}({s.get('score')})" for s in src[:3]
                )
            except json.JSONDecodeError:
                sources = q.sources or ""
            w.writerow(
                [
                    q.row_number,
                    q.section,
                    q.question_text,
                    q.answer_text,
                    q.confidence,
                    q.status,
                    sources,
                    q.notes,
                ]
            )
    qn.status = "exported"
    db.commit()
    return path


def export_xlsx(db: Session, questionnaire_id: int) -> Path:
    qn = db.get(Questionnaire, questionnaire_id)
    if not qn:
        raise ValueError("Not found")
    questions = (
        db.query(Question)
        .filter(Question.questionnaire_id == questionnaire_id)
        .order_by(Question.row_number)
        .all()
    )
    wb = Workbook()
    ws = wb.active
    ws.title = "Answers"
    headers = [
        "Row",
        "Section",
        "Question",
        "Answer",
        "Confidence",
        "Status",
        "Top Source",
        "Notes",
    ]
    header_fill = PatternFill("solid", fgColor="1F2937")
    header_font = Font(color="FFFFFF", bold=True)
    for col, h in enumerate(headers, 1):
        cell = ws.cell(1, col, h)
        cell.fill = header_fill
        cell.font = header_font

    low_fill = PatternFill("solid", fgColor="FEE2E2")
    mid_fill = PatternFill("solid", fgColor="FEF3C7")
    high_fill = PatternFill("solid", fgColor="D1FAE5")

    for i, q in enumerate(questions, 2):
        top_source = ""
        try:
            src = json.loads(q.sources or "[]")
            if src:
                top_source = f"{src[0].get('title')}: {str(src[0].get('excerpt', ''))[:120]}"
        except json.JSONDecodeError:
            pass
        values = [
            q.row_number,
            q.section,
            q.question_text,
            q.answer_text,
            q.confidence,
            q.status,
            top_source,
            q.notes,
        ]
        for col, val in enumerate(values, 1):
            cell = ws.cell(i, col, val)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        conf = q.confidence or 0
        fill = high_fill if conf >= 0.7 else mid_fill if conf >= 0.4 else low_fill
        ws.cell(i, 5).fill = fill

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 45
    ws.column_dimensions["D"].width = 55
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 12
    ws.column_dimensions["G"].width = 40
    ws.column_dimensions["H"].width = 20

    path = EXPORT_DIR / f"questionnaire_{questionnaire_id}_answers.xlsx"
    wb.save(path)
    qn.status = "exported"
    db.commit()
    return path
