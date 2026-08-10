"""Extract plain text from uploaded policy / document files."""

from pathlib import Path

from pypdf import PdfReader


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix in {".txt", ".md", ".markdown", ".csv"}:
        return path.read_text(encoding="utf-8", errors="ignore")
    if suffix == ".docx":
        return _extract_docx(path)
    return path.read_text(encoding="utf-8", errors="ignore")


def _extract_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    parts: list[str] = []
    for page in reader.pages:
        text = page.extract_text() or ""
        if text.strip():
            parts.append(text)
    return "\n\n".join(parts)


def _extract_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as e:
        raise ValueError("python-docx is required for DOCX support") from e
    doc = Document(str(path))
    parts: list[str] = []
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t:
            parts.append(t)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def chunk_text(text: str, chunk_size: int = 900, overlap: int = 120) -> list[str]:
    text = " ".join(text.split())
    if not text:
        return []
    chunks: list[str] = []
    start = 0
    n = len(text)
    while start < n:
        end = min(start + chunk_size, n)
        if end < n:
            window = text[start:end]
            for sep in (". ", "。", "\n"):
                idx = window.rfind(sep)
                if idx > chunk_size * 0.4:
                    end = start + idx + len(sep)
                    break
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= n:
            break
        start = max(end - overlap, start + 1)
    return chunks
